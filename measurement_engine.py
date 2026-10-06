# -*- coding: utf-8 -*-
"""
Единый движок измерений.
Один цикл по частотам, внутри которого выполняются ВСЕ выбранные тесты.

Активные тесты передаются списком, например:
    ['freq', 'level']
    ['freq', 'level', 'am_noise', 'fm_noise', 'pm_noise']

Нереализованные тесты (am_noise, fm_noise, pm_noise) — пока пропускаются
с сообщением в лог.
"""

import time
import os
from datetime import datetime


# ============================================================
#                    КОНСТАНТЫ
# ============================================================

# Список частот для теста (МГц)
F_LIST = [
    '0.01', '0.1', '1', '10', '50', '100', '250', '500',
    '1000', '1500', '2000', '2500', '3000'
]

# Уровни для теста уровня (dBm):
U_LIST_0 = [-13.0, 0.0]           # для частот 0.01 и 0.1 МГц
U_LIST   = [-13.0, 0.0, 13.0]     # для остальных

# Допуски
FREQ_TOLERANCE_FORMULA = lambda f_hz: 5e-6 * abs(f_hz) + 1.0   # Гц
LEVEL_TOLERANCE_DB = 1.0                                        # дБ

# Порог «жёлтой зоны» — доля от допуска
TOLERANCE_RATIO_WARN = 0.7

# Задержка после установки параметров (сек) — везде одинаковая
SETTLE_DELAY = 6.0

# Начальный уровень генератора (dBm)
INITIAL_POWER_DBM = -13

# Идентификаторы активных тестов
TEST_FREQ = 'freq'
TEST_LEVEL = 'level'
TEST_AM_NOISE = 'am_noise'
TEST_FM_NOISE = 'fm_noise'
TEST_PM_NOISE = 'pm_noise'

# Тесты, которые пока НЕ реализованы — движок их пропустит
NOT_IMPLEMENTED = {TEST_AM_NOISE, TEST_FM_NOISE, TEST_PM_NOISE}


# ============================================================
#                    ЦВЕТА / ОЦЕНКИ
# ============================================================

def evaluate(value, tolerance):
    """
    Оценка результата по допуску.
    Возвращает (status_code, comment_emoji):
      'ok'   — 🟢
      'warn' — 🟡
      'fail' — 🔴
    """
    abs_v = abs(value)
    if abs_v > tolerance:
        return 'fail', '🔴'
    elif abs_v > TOLERANCE_RATIO_WARN * tolerance:
        return 'warn', '🟡'
    else:
        return 'ok', '🟢'


# ============================================================
#                    ОСНОВНОЙ ДВИЖОК
# ============================================================

class MeasurementEngine:
    """
    Единый движок. Один цикл по частотам.
    Внутри каждой частоты — выбранные измерения.
    """

    def __init__(self, gen, sk3, active_tests=None, log_callback=None):
        """
        :param gen: объект pyvisa-сессии генератора SMB100A
        :param sk3: объект serial.Serial для СК3-50
        :param active_tests: список идентификаторов активных тестов.
                             По умолчанию — все реализованные (freq, level).
        :param log_callback: функция f(msg: str)
        """
        self.gen = gen
        self.sk3 = sk3
        self.log = log_callback if log_callback is not None else (lambda msg: None)

        if active_tests is None:
            active_tests = [TEST_FREQ, TEST_LEVEL]
        self.active_tests = list(active_tests)

        # Результаты по каждому тесту — свой список
        self.results = {
            TEST_FREQ: [],
            TEST_LEVEL: [],
            TEST_AM_NOISE: [],
            TEST_FM_NOISE: [],
            TEST_PM_NOISE: [],
        }

        # Общий итоговый статус (done / warn / error)
        self.final_status = 'not_started'

        # Признак, что где-то встретился warn / fail
        self._any_warn = False
        self._stop_requested = False

    # --------------------------------------------------------
    #                    ИНИЦИАЛИЗАЦИЯ
    # --------------------------------------------------------

    def _init_generator(self):
        self.log("Инициализация генератора SMB100A…")
        try:
            self.gen.write('*RST')
            self.gen.write('OUTP ON')
            self.gen.write('FREQ 10 MHz')
            self.gen.write(f'POW {INITIAL_POWER_DBM}')
            time.sleep(1)
            self.log("Генератор: *RST, OUTP ON, FREQ 10 MHz, POW -13 dBm.")
        except Exception as e:
            self.log(f"⚠ Ошибка инициализации генератора: {e}")
            raise

    def _deinit_generator(self):
        self.log("Выключение выхода генератора (OUTP OFF)…")
        try:
            self.gen.write('OUTP OFF')
        except Exception as e:
            self.log(f"⚠ Не удалось выключить выход генератора: {e}")

    # --------------------------------------------------------
    #                    ПОДСЧЁТ ТОЧЕК
    # --------------------------------------------------------

    def _count_total_points(self, freqs):
        """Сколько всего измерений будет сделано (для прогресс-бара)."""
        total = 0
        for f_mhz in freqs:
            f_val = float(f_mhz)
            if TEST_FREQ in self.active_tests:
                total += 1
            if TEST_LEVEL in self.active_tests:
                levels = U_LIST_0 if f_val in (0.01, 0.1) else U_LIST
                total += len(levels)
            # нереализованные не считаем
        return total

    # --------------------------------------------------------
    #                    ОСНОВНОЙ ЗАПУСК
    # --------------------------------------------------------

    def run(self, freqs=None, progress_callback=None):
        """
        :param freqs: список частот (МГц). По умолчанию F_LIST.
        :param progress_callback: f(idx, total, info_str) — перед каждой точкой.
                                  Может возбудить InterruptedError для остановки.
        :return: (results_dict, final_status)
        """
        if freqs is None:
            freqs = F_LIST

        self.results = {k: [] for k in self.results}
        self.final_status = 'not_started'
        self._any_warn = False
        self._stop_requested = False

        total_points = self._count_total_points(freqs)

        self.log("─" * 50)
        self.log(f"Активные тесты: {', '.join(self.active_tests)}")
        self.log(f"Частот в цикле: {len(freqs)}. Всего измерений: {total_points}")

        # Проверяем, не выбраны ли только нереализованные тесты
        implemented = [t for t in self.active_tests if t not in NOT_IMPLEMENTED]
        if not implemented:
            self.log("❌ Не выбрано ни одного реализованного теста.")
            self.final_status = 'error'
            return self.results, self.final_status

        # Инициализация
        try:
            self._init_generator()
        except Exception as e:
            self.log(f"❌ Не удалось инициализировать генератор: {e}")
            self.final_status = 'error'
            return self.results, self.final_status

        point_idx = 0

        try:
            for freq_i, freq_mhz in enumerate(freqs, start=1):
                self.log(f"─── Частота {freq_mhz} МГц ({freq_i}/{len(freqs)}) ───")

                # Устанавливаем частоту
                try:
                    self.gen.write(f"FREQ {freq_mhz} MHz")
                except Exception as e:
                    self.log(f"❌ Ошибка установки частоты {freq_mhz} МГц: {e}")
                    self.final_status = 'error'
                    return self.results, self.final_status

                f_val = float(freq_mhz)

                # --- Тест частоты ---
                if TEST_FREQ in self.active_tests:
                    point_idx += 1
                    info = f"Частота {freq_mhz} МГц — измерение частоты"
                    if progress_callback:
                        progress_callback(point_idx, total_points, info)

                    row, status = self._measure_freq(freq_mhz)
                    if row is not None:
                        self.results[TEST_FREQ].append(row)
                    if status == 'fail':
                        self.log("🔴 ПРОВАЛ на тесте частоты — остановка.")
                        self.final_status = 'error'
                        return self.results, self.final_status
                    if status == 'warn':
                        self._any_warn = True

                # --- Тест уровня ---
                if TEST_LEVEL in self.active_tests:
                    levels = U_LIST_0 if f_val in (0.01, 0.1) else U_LIST
                    for lv in levels:
                        point_idx += 1
                        info = f"Частота {freq_mhz} МГц — уровень {lv:+.0f} dBm"
                        if progress_callback:
                            progress_callback(point_idx, total_points, info)

                        row, status = self._measure_level(freq_mhz, lv)
                        if row is not None:
                            self.results[TEST_LEVEL].append(row)
                        if status == 'fail':
                            self.log("🔴 ПРОВАЛ на тесте уровня — остановка.")
                            self.final_status = 'error'
                            return self.results, self.final_status
                        if status == 'warn':
                            self._any_warn = True

                # --- Нереализованные тесты (пока пропускаем) ---
                for t in self.active_tests:
                    if t in NOT_IMPLEMENTED:
                        self.log(f"   (тест «{t}» пока в разработке — пропущен)")

            # Все частоты пройдены
            if self._any_warn:
                self.final_status = 'warn'
                self.log("🟡 Тест завершён. Есть предупреждения.")
            else:
                self.final_status = 'done'
                self.log("🟢 Тест завершён успешно. Все точки в норме.")

            return self.results, self.final_status

        finally:
            self._deinit_generator()

    # --------------------------------------------------------
    #                    ИЗМЕРЕНИЕ ЧАСТОТЫ
    # --------------------------------------------------------

    def _measure_freq(self, freq_mhz):
        """
        Измерение частоты на одной точке.
        Возвращает: (row, status)
        row = [freq_set_hz, freq_meas_hz, error_hz, tolerance_hz,
               if_khz, status_code, comment_emoji]
        """
        # Пауза — чтобы прибор успел установиться
        time.sleep(SETTLE_DELAY)

        try:
            # Установленная частота на генераторе (двойной запрос)
            _ = float(self.gen.query("FREQ?")) / 1_000_000
            inpf1 = float(self.gen.query("FREQ?"))

            # Измеренная частота СК3-50
            self.sk3.write(b'FREQ?\n')
            self.sk3.flush()
            raw = self.sk3.readline()
            fiz = float(raw.decode('utf-8', errors='ignore').strip())

            # Погрешность и допуск
            kf = fiz - inpf1
            tolerance = FREQ_TOLERANCE_FORMULA(inpf1)

            # ПЧ
            self.sk3.write(b'FREQ:INT?\n')
            self.sk3.flush()
            raw_if = self.sk3.readline()
            try:
                fph = float(raw_if.decode('utf-8', errors='ignore').strip()) / 1000.0
            except ValueError:
                fph = float('nan')

            status, emoji = evaluate(kf, tolerance)

            self.log(f"   freq : {kf:+.2f} Гц, допуск ±{tolerance:.2f} Гц → {emoji}")

            row = [inpf1, fiz, round(kf, 2), round(tolerance, 2),
                   round(fph, 2), status, emoji]
            return row, status

        except Exception as e:
            self.log(f"   ❌ Ошибка измерения частоты: {e}")
            return None, 'error'

    # --------------------------------------------------------
    #                    ИЗМЕРЕНИЕ УРОВНЯ
    # --------------------------------------------------------

    def _measure_level(self, freq_mhz, level_dbm):
        """
        Измерение уровня на одной точке.
        row = [freq_mhz, level_set_dbm, level_meas_dbm,
               error_db, tolerance_db, status_code, comment_emoji]
        """
        try:
            self.gen.write(f"POW {level_dbm}")
        except Exception as e:
            self.log(f"   ❌ Ошибка установки уровня {level_dbm} dBm: {e}")
            return None, 'error'

        time.sleep(SETTLE_DELAY)

        try:
            self.sk3.write(b'POW?\n')
            self.sk3.flush()
            raw = self.sk3.readline()
            oup = float(raw.decode('utf-8', errors='ignore').strip())

            ku = oup - level_dbm
            status, emoji = evaluate(ku, LEVEL_TOLERANCE_DB)

            self.log(f"   level: {ku:+.2f} дБ, допуск ±{LEVEL_TOLERANCE_DB:.2f} дБ → {emoji}")

            row = [float(freq_mhz), level_dbm, round(oup, 2),
                   round(ku, 2), LEVEL_TOLERANCE_DB, status, emoji]
            return row, status

        except Exception as e:
            self.log(f"   ❌ Ошибка измерения уровня: {e}")
            return None, 'error'

    # --------------------------------------------------------
    #                    СОХРАНЕНИЕ
    # --------------------------------------------------------

    def save_to_excel(self, filepath):
        """
        Сохраняет результаты в Excel.
        Каждый непустой набор — на отдельный лист.
        С цветной заливкой колонки «Статус».
        """
        import pandas as pd
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill

        # Собираем листы
        sheets = {}

        if self.results[TEST_FREQ]:
            sheets['Частота'] = pd.DataFrame(
                self.results[TEST_FREQ],
                columns=['Уст Частота, Гц', 'Изм Частота, Гц', 'Погрешность, Гц',
                         'Допуск, Гц', 'Частота ПЧ, кГц', 'Статус', 'Комментарий'])

        if self.results[TEST_LEVEL]:
            sheets['Уровень'] = pd.DataFrame(
                self.results[TEST_LEVEL],
                columns=['Частота, МГц', 'Заданный уровень, дБм', 'Измеренный уровень, дБм',
                         'Погрешность, дБ', 'Допуск, дБ', 'Статус', 'Комментарий'])

        if not sheets:
            self.log("⚠ Нет данных для сохранения.")
            return

        # Определяем режим
        if os.path.exists(filepath):
            kwargs = {'engine': 'openpyxl', 'mode': 'a', 'if_sheet_exists': 'replace'}
        else:
            kwargs = {'engine': 'openpyxl', 'mode': 'w'}

        with pd.ExcelWriter(filepath, **kwargs) as writer:
            for name, df in sheets.items():
                # Перед записью: заменяем в колонке «Статус» код на пустую строку
                df_to_save = df.copy()
                df_to_save['Статус'] = ""   # пусто, зальём цветом
                df_to_save['Комментарий'] = df['Комментарий'].apply(
                    lambda e: e if isinstance(e, str) else "")
                df_to_save.to_excel(writer, sheet_name=name, index=False)

        # Дальше — раскрашиваем ячейки через openpyxl
        wb = load_workbook(filepath)

        fills = {
            'ok':   PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid'),
            'warn': PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid'),
            'fail': PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid'),
        }

        for sheet_name, df in sheets.items():
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            # Колонка «Статус» — шестая (F), начиная со 2-й строки
            col_letter = 'F'
            for i, status_code in enumerate(df['Статус'], start=2):
                cell = ws[f'{col_letter}{i}']
                if status_code in fills:
                    cell.fill = fills[status_code]

        wb.save(filepath)
        self.log(f"Результаты сохранены: {filepath}")

    # --------------------------------------------------------
    #                    СВОДКА
    # --------------------------------------------------------

    def summary(self):
        """Возвращает dict со сводкой по каждому тесту."""
        out = {}
        for key, rows in self.results.items():
            if not rows:
                continue
            total = len(rows)
            ok = sum(1 for r in rows if r[5] == 'ok')
            warn = sum(1 for r in rows if r[5] == 'warn')
            fail = sum(1 for r in rows if r[5] == 'fail')
            out[key] = {'total': total, 'ok': ok, 'warn': warn, 'fail': fail}
        out['final_status'] = self.final_status
        return out