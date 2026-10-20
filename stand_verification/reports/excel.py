import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill


STATUS_LABELS = {
    'ok':    'Пройден',
    'warn':  'Внимание',
    'fail':  'Провален',
    'error': 'Ошибка',
}


def _fill_for(status):
    return {
        'ok':    'C6EFCE',
        'warn':  'FFEB9C',
        'fail':  'FFC7CE',
        'error': 'FFC7CE',
    }.get(status)


def export_excel(snapshot, filepath, log_callback=None):
    if os.path.exists(filepath):
        wb = load_workbook(filepath)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    layouts = {
        'freq': (
            'Частота',
            ['Уст Частота, Гц', 'Изм Частота, Гц', 'Погрешность, Гц',
             'Допуск, Гц', 'Частота ПЧ, кГц', 'Статус'],
        ),
        'level': (
            'Уровень',
            ['Частота, МГц', 'Заданный уровень, дБм', 'Измеренный уровень, дБм',
             'Погрешность, дБ', 'Допуск, дБ', 'Статус'],
        ),
        'am_noise': (
            'АМ шум',
            ['Уст Частота, Гц', 'Полоса, кГц', 'СКЗ шум, %', 'Допуск, %', 'Статус'],
        ),
        'fm_noise': (
            'ЧМ шум',
            ['Уст Частота, Гц', 'Полоса, кГц', 'СКЗ шум, Гц', 'Допуск, Гц', 'Статус'],
        ),
        'pm_noise': (
            'ФМ шум',
            ['Уст Частота, Гц', 'Полоса, кГц', 'СКЗ шум, рад', 'Допуск, рад', 'Статус'],
        ),
        'am_fm_transition': (
            'Переход АМ↔ЧМ',
            ['Направление', 'Значение +', 'Значение −',
             'Коэффициент', 'Допуск', 'Ед.', 'Статус'],
        ),
        'thd_am': (
            'КГ АМ',
            ['Fm, Гц', 'Заданный Кам, %', 'Измеренный Кам, %',
             'КГ СК3-50, %', 'КГ С6-22, %', 'Допуск, %',
             'Изм. Частота С6-22', 'Фильтр, кГц', 'Статус'],
        ),
        'thd_fm': (
            'КГ ЧМ',
            ['Fm, Гц', 'Заданная девиация, Гц', 'Измеренная девиация, Гц',
             'КГ СК3-50, %', 'КГ С6-22, %', 'Допуск, %',
             'Изм. Частота С6-22', 'Фильтр, кГц', 'Статус'],
        ),
        'pm_peaks': (
            'PM',
            ['Fm, Гц', 'Индекс заданный', '+P, рад', '-P, рад',
             '+dP, рад', '-dP, рад', 'Допуск, рад',
             'Фильтр, кГц', 'Статус'],
        ),
        'am_peaks': (
            'АМ',
            ['F несущая, МГц', 'Fm, Гц', 'Кам заданный, %',
             '+Кам, %', '-Кам, %', '+dКам, %', '-dКам, %',
             'Допуск, %', 'Фильтр, кГц', 'Статус'],
        ),
        'fm_linearity': (
            'ЧМ_лин',
            ['Шкала, кГц', 'Дев част, кГц', 'Делитель',
             '+F, кГц', '-F, кГц',
             '+dF, кГц', '-dF, кГц', 'Допуск, кГц', 'Статус'],
        ),
        'fm_deviation': (
            'ЧМ',
            ['Fm, кГц', 'Дев част, кГц', '+F, кГц', '-F, кГц',
             '+dF, кГц', '-dF, кГц', 'Допуск, кГц',
             'Фильтр, кГц', 'Статус'],
        ),
    }

    for name, _ in layouts.values():
        if name in wb.sheetnames:
            del wb[name]

    for key, (name, headers) in layouts.items():
        if key not in snapshot.active_tests:
            continue

        ws = wb.create_sheet(name)
        ws.append(headers)

        for row in snapshot.results.get(key, []):
            if key == 'freq':
                values = [
                    row.reference_frequency_hz, row.measured_frequency_hz,
                    row.error_hz, row.tolerance_hz, row.if_frequency_khz,
                    STATUS_LABELS.get(row.status, row.status),
                ]
                status_col = 6
            elif key == 'level':
                values = [
                    row.frequency_mhz, row.set_level_dbm, row.measured_level_dbm,
                    row.error_db, row.tolerance_db,
                    STATUS_LABELS.get(row.status, row.status),
                ]
                status_col = 6
            elif key in ('am_noise', 'fm_noise', 'pm_noise'):
                band_str = f"{row.band_low_khz:g}-{row.band_high_khz:g}"
                values = [
                    row.frequency_mhz * 1_000_000,
                    band_str,
                    row.noise_value, row.tolerance,
                    STATUS_LABELS.get(row.status, row.status),
                ]
                status_col = 5
            elif key == 'am_fm_transition':
                values = [
                    row.direction, row.value_plus, row.value_minus,
                    row.coefficient, row.tolerance, row.unit,
                    STATUS_LABELS.get(row.status, row.status),
                ]
                status_col = 7
            elif key == 'thd_am':
                values = [
                    row.mod_freq_hz, row.set_depth, row.meas_depth,
                    row.thd_sk3, row.thd_c6, row.tolerance,
                    row.mod_freq_c6, row.band_str,
                    STATUS_LABELS.get(row.status, row.status),
                ]
                status_col = 9
            elif key == 'thd_fm':
                values = [
                    row.mod_freq_hz, row.set_deviation_hz, row.meas_deviation_hz,
                    row.thd_sk3, row.thd_c6, row.tolerance,
                    row.mod_freq_c6, row.band_str,
                    STATUS_LABELS.get(row.status, row.status),
                ]
                status_col = 9
            elif key == 'pm_peaks':
                values = [
                    row.mod_freq_hz, row.set_index,
                    row.peak_plus_rad, row.peak_minus_rad,
                    row.delta_plus_rad, row.delta_minus_rad,
                    row.tolerance_rad, row.band_str,
                    STATUS_LABELS.get(row.status, row.status),
                ]
                status_col = 9
            elif key == 'am_peaks':
                values = [
                    row.carrier_mhz, row.mod_freq_hz, row.set_depth_percent,
                    row.peak_plus_percent, row.peak_minus_percent,
                    row.delta_plus_percent, row.delta_minus_percent,
                    row.tolerance_percent, row.band_str,
                    STATUS_LABELS.get(row.status, row.status),
                ]
                status_col = 10
        
            elif key == 'fm_linearity':
                values = [
                    row.scale_khz, row.set_deviation_khz, row.divisor,
                    row.peak_plus_khz, row.peak_minus_khz,
                    row.delta_plus_khz, row.delta_minus_khz,
                    row.tolerance_khz,
                    STATUS_LABELS.get(row.status, row.status),
                ]
                status_col = 9
            elif key == 'fm_deviation':
                values = [
                    row.mod_freq_khz, row.set_deviation_khz,
                    row.peak_plus_khz, row.peak_minus_khz,
                    row.delta_plus_khz, row.delta_minus_khz,
                    row.tolerance_khz, row.band_str,
                    STATUS_LABELS.get(row.status, row.status),
                ]
                status_col = 9
            else:
                continue

            ws.append(values)

            color = _fill_for(row.status)
            if color:
                fill = PatternFill(start_color=color, end_color=color,
                                   fill_type='solid')
                ws.cell(ws.max_row, status_col).fill = fill

    # --- Служебные листы ---
    for name in ('Сводка', 'Тесты', 'Диагностика'):
        if name in wb.sheetnames:
            del wb[name]

    ws = wb.create_sheet('Сводка')
    ws.append(['Итог', 'Причина завершения', 'Выключение выполнено',
               'Ошибка выключения'])
    ws.append([
        snapshot.final_status, snapshot.termination_reason,
        snapshot.shutdown_success, snapshot.shutdown_error,
    ])

    ws = wb.create_sheet('Тесты')
    ws.append(['Тест', 'Статус', 'План', 'Точек', 'OK',
               'Предупреждений', 'Ошибок'])
    for key, value in snapshot.summary.items():
        if key == 'final_status':
            continue
        ws.append([
            key, value['status'], value['planned'], value['total'],
            value['ok'], value['warn'], value['fail'] + value['error'],
        ])

    ws = wb.create_sheet('Диагностика')
    ws.append(['Диагностика'])
    for diagnostic in snapshot.diagnostics:
        ws.append([str(diagnostic)])

    wb.save(filepath)
    if log_callback:
        log_callback(f'Результаты сохранены: {filepath}')
    return True
