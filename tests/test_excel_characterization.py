import importlib.util
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from stand_verification.measurement.engine import MeasurementEngine, TEST_FREQ, TEST_LEVEL
from stand_verification.models import FrequencyResult, LevelResult
from stand_verification.reports.excel import export_excel


@unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'openpyxl is not installed')
class ExcelCharacterizationTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'report.xlsx'
        self.logs = []
        self.engine = MeasurementEngine(None, None, [TEST_FREQ, TEST_LEVEL],
                                        log_callback=self.logs.append)

    def save_and_load(self):
        from openpyxl import load_workbook
        self.assertIs(export_excel(self.engine.snapshot(), self.path, self.engine._log), True)
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        return workbook

    def test_layout_values_colors_comments_and_diagnostics(self):
        self.engine.results[TEST_FREQ] = [
            FrequencyResult(
                reference_frequency_hz=10000, measured_frequency_hz=10000,
                error_hz=0, tolerance_hz=1.05, if_frequency_khz=2,
                status='ok', comment='accepted',
            ),
            FrequencyResult(
                reference_frequency_hz=10000, measured_frequency_hz=10001,
                error_hz=1, tolerance_hz=1.05, if_frequency_khz=2,
                status='warn', comment='near limit',
            ),
            FrequencyResult(
                reference_frequency_hz=10000, measured_frequency_hz=10002,
                error_hz=2, tolerance_hz=1.05, if_frequency_khz=2,
                status='fail', comment='outside limit',
            ),
            FrequencyResult(
                reference_frequency_hz=10000, measured_frequency_hz=None,
                error_hz=None, tolerance_hz=None, if_frequency_khz=None,
                status='error', comment='meter failure',
            ),
        ]
        self.engine.results[TEST_LEVEL] = [LevelResult(
            frequency_mhz=2000, set_level_dbm=0, measured_level_dbm=1.5,
            error_db=1.5, tolerance_db=2, status='warn', comment='near limit',
        )]
        self.engine._planned = {TEST_FREQ: 5, TEST_LEVEL: 1}
        self.engine.final_status = 'error'
        self.engine.termination_reason = 'meter failure'
        self.engine.shutdown_success = False
        self.engine.shutdown_error = 'output failure'
        self.engine.diagnostics = ['meter failure', 'output failure']
        workbook = self.save_and_load()
        self.assertEqual(workbook.sheetnames,
                         ['Частота', 'Уровень', 'Сводка', 'Тесты', 'Диагностика'])
        frequency = workbook['Частота']
        self.assertEqual(list(frequency.values)[0], (
            'Уст Частота, Гц', 'Изм Частота, Гц', 'Погрешность, Гц',
            'Допуск, Гц', 'Частота ПЧ, кГц', 'Статус', 'Комментарий'))
        self.assertEqual(list(frequency.values)[1:], [
            (10000, 10000, 0, 1.05, 2, 'ok', 'accepted'),
            (10000, 10001, 1, 1.05, 2, 'warn', 'near limit'),
            (10000, 10002, 2, 1.05, 2, 'fail', 'outside limit'),
            (10000, None, None, None, None, 'error', 'meter failure'),
        ])
        for row, color in zip(range(2, 6), ['C6EFCE', 'FFEB9C', 'FFC7CE', 'FFC7CE']):
            with self.subTest(row=row):
                self.assertEqual(frequency.cell(row, 6).fill.fill_type, 'solid')
                self.assertEqual(frequency.cell(row, 6).fill.fgColor.rgb, '00' + color)
        self.assertIsNone(frequency['G2'].comment)
        self.assertIsNone(frequency['G3'].comment)
        self.assertIsNone(frequency['G4'].comment)
        self.assertEqual(frequency['G5'].comment.text, 'meter failure')
        self.assertEqual(frequency['G5'].comment.author, 'MeasurementEngine')
        self.engine.results[TEST_LEVEL].append(LevelResult(
            frequency_mhz=2000, set_level_dbm=0, measured_level_dbm=None,
            error_db=None, tolerance_db=2, status='error', comment='meter failure',
        ))
        self.engine._planned[TEST_LEVEL] = 2
        workbook = self.save_and_load()
        self.assertEqual(list(workbook['Уровень'].values), [
            ('Частота, МГц', 'Заданный уровень, дБм', 'Измеренный уровень, дБм',
             'Погрешность, дБ', 'Допуск, дБ', 'Статус', 'Комментарий'),
            (2000, 0, 1.5, 1.5, 2, 'warn', 'near limit'),
            (2000, 0, None, None, 2, 'error', 'meter failure'),
        ])
        self.assertEqual(workbook['Уровень']['F3'].fill.fgColor.rgb, '00FFC7CE')
        self.assertEqual(workbook['Уровень']['G3'].comment.text, 'meter failure')
        self.assertEqual(workbook['Уровень']['G3'].comment.author, 'MeasurementEngine')
        self.assertEqual(list(workbook['Сводка'].values), [
            ('Итог', 'Причина завершения', 'Выключение выполнено', 'Ошибка выключения'),
            ('error', 'meter failure', False, 'output failure'),
        ])
        self.assertEqual(list(workbook['Тесты'].values), [
            ('Тест', 'Статус', 'План', 'Точек', 'OK', 'Предупреждений', 'Ошибок'),
            ('freq', 'error', 5, 4, 1, 1, 2),
            ('level', 'error', 2, 2, 0, 1, 1),
        ])
        self.assertEqual(list(workbook['Диагностика'].values),
                         [('Диагностика',), ('meter failure',), ('output failure',)])
        self.assertEqual(self.logs[-1], f'Результаты сохранены: {self.path}')

    def test_selected_empty_test_has_header_and_unselected_test_is_absent(self):
        self.engine.active_tests = [TEST_FREQ]
        workbook = self.save_and_load()
        self.assertIn('Частота', workbook.sheetnames)
        self.assertNotIn('Уровень', workbook.sheetnames)
        self.assertEqual(workbook['Частота'].max_row, 1)
        self.assertEqual(list(workbook['Тесты'].values)[1],
                         ('freq', 'not_started', 0, 0, 0, 0, 0))

    def test_resave_replaces_managed_sheets_and_preserves_unrelated_sheet(self):
        from openpyxl import Workbook
        original = Workbook()
        original.active.title = 'User data'
        original.active['A1'] = 'preserve me'
        for name in ('Частота', 'Уровень', 'Сводка', 'Тесты', 'Диагностика'):
            original.create_sheet(name)['A1'] = 'stale'
        original.save(self.path)
        original.close()
        self.engine.active_tests = [TEST_LEVEL]
        self.engine.results[TEST_LEVEL] = [LevelResult(
            frequency_mhz=1, set_level_dbm=0, measured_level_dbm=0,
            error_db=0, tolerance_db=1, status='ok', comment='accepted',
        )]
        first = self.save_and_load()
        self.assertNotIn('Частота', first.sheetnames)
        self.assertEqual(first['User data']['A1'].value, 'preserve me')
        self.assertEqual(first['Уровень'].max_row, 2)
        first.close()
        self.engine.results[TEST_LEVEL] = []
        second = self.save_and_load()
        self.assertEqual(second.sheetnames,
                         ['User data', 'Уровень', 'Сводка', 'Тесты', 'Диагностика'])
        self.assertEqual(second['User data']['A1'].value, 'preserve me')
        self.assertEqual(second['Уровень'].max_row, 1)

    def test_failed_save_does_not_log_success_or_change_results(self):
        self.engine.results[TEST_LEVEL] = [LevelResult(
            frequency_mhz=1, set_level_dbm=0, measured_level_dbm=0,
            error_db=0, tolerance_db=1, status='ok', comment='accepted',
        )]
        with self.assertRaises(OSError):
            export_excel(self.engine.snapshot(), self.path / 'missing' / 'report.xlsx',
                         self.engine._log)
        self.assertEqual(self.logs, [])
        self.assertEqual(self.engine.results[TEST_LEVEL],
                         [LevelResult(
                             frequency_mhz=1, set_level_dbm=0,
                             measured_level_dbm=0, error_db=0, tolerance_db=1,
                             status='ok', comment='accepted',
                         )])

    def test_snapshot_is_independent_of_later_engine_mutations(self):
        self.engine.results[TEST_FREQ] = [FrequencyResult(
            reference_frequency_hz=10000, measured_frequency_hz=10000,
            error_hz=0, tolerance_hz=1.05, if_frequency_khz=2,
            status='ok', comment='accepted',
        )]
        self.engine.final_status = 'done'
        snapshot = self.engine.snapshot()
        self.engine.results[TEST_FREQ][0] = replace(self.engine.results[TEST_FREQ][0], measured_frequency_hz=-1)
        self.engine.results[TEST_FREQ].append(FrequencyResult(
            reference_frequency_hz=20000, measured_frequency_hz=20000,
            error_hz=0, tolerance_hz=1.1, if_frequency_khz=2,
            status='ok', comment='accepted',
        ))
        self.engine.active_tests.clear()
        self.engine.diagnostics.append('late diagnostic')
        self.assertEqual(snapshot.results[TEST_FREQ][0].measured_frequency_hz, 10000)
        self.assertEqual(len(snapshot.results[TEST_FREQ]), 1)
        self.assertEqual(snapshot.active_tests, (TEST_FREQ, TEST_LEVEL))
        self.assertEqual(snapshot.diagnostics, ())
        self.engine.results.clear()
        self.engine.final_status = 'error'
        self.assertEqual(snapshot.results[TEST_FREQ][0].measured_frequency_hz, 10000)
        self.assertEqual(snapshot.final_status, 'done')
        self.assertEqual(snapshot.summary[TEST_FREQ]['status'], 'incomplete')
        self.assertEqual(snapshot.summary[TEST_FREQ]['total'], 1)

    def test_snapshot_deep_copies_nested_summary(self):
        summary = {TEST_FREQ: {'status': 'done', 'total': 1}}
        with patch.object(self.engine, 'summary', return_value=summary):
            snapshot = self.engine.snapshot()
        summary[TEST_FREQ]['status'] = 'error'
        summary[TEST_FREQ]['total'] = 0
        self.assertEqual(snapshot.summary[TEST_FREQ], {'status': 'done', 'total': 1})
