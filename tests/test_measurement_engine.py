import tempfile
import unittest
import importlib.util
from pathlib import Path
from unittest.mock import patch

from stand_verification.measurement.engine import MeasurementEngine, TEST_FREQ, TEST_LEVEL
from stand_verification.models import FrequencyResult, LevelResult
from stand_verification.reports.excel import export_excel


class FakeGenerator:
    def __init__(self, fail_write=None):
        self.commands = []
        self.fail_write = fail_write

    def write(self, command):
        self.commands.append(command)
        if command == self.fail_write:
            raise OSError('write failed: ' + command)

    def query(self, command):
        return '10000'


class FakeMeter:
    def __init__(self, responses=None):
        self.responses = list(responses or [b'10000\n', b'2000\n', b'0\n'])
        self.request = b''

    def write(self, command):
        self.request = command

    def flush(self):
        pass

    def readline(self):
        if not self.responses:
            raise OSError('no response')
        return self.responses.pop(0)


class MeasurementEngineTests(unittest.TestCase):
    def engine(self, gen=None, sk3=None, tests=None, stop_event=None):
        return MeasurementEngine(gen or FakeGenerator(), sk3 or FakeMeter(), tests or [TEST_FREQ], stop_event=stop_event)

    def test_success_has_cleanup_and_summary(self):
        engine = self.engine()
        with patch('stand_verification.measurement.frequency.SETTLE_DELAY', 0):
            results, status = engine.run(['0.01'])
        self.assertEqual(status, 'done')
        self.assertEqual(len(results[TEST_FREQ]), 1)
        self.assertEqual(engine.shutdown_success, True)
        self.assertEqual(engine.summary()[TEST_FREQ]['status'], 'done')
        self.assertIn('OUTP OFF', engine.gen.session.commands)

    def test_rerun_does_not_mutate_prior_snapshot(self):
        engine = self.engine(sk3=FakeMeter([
            b'10000\n', b'2000\n', b'10001\n', b'2000\n',
        ]))
        with patch('stand_verification.measurement.frequency.SETTLE_DELAY', 0):
            engine.run(['0.01'])
            snapshot = engine.snapshot()
            engine.run(['0.01'])
        self.assertEqual(snapshot.results[TEST_FREQ], [FrequencyResult(
            reference_frequency_hz=10000.0, measured_frequency_hz=10000.0,
            error_hz=0.0, tolerance_hz=1.05, if_frequency_khz=2.0,
            status='ok', comment='\U0001f7e2',
        )])
        self.assertEqual(len(snapshot.results[TEST_FREQ]), 1)
        self.assertEqual(snapshot.summary[TEST_FREQ]['total'], 1)
        self.assertEqual(engine.results[TEST_FREQ][0].measured_frequency_hz, 10001.0)

    def test_invalid_if_records_error_and_stops(self):
        engine = self.engine(sk3=FakeMeter([b'10000\n', b'invalid\n']))
        with patch('stand_verification.measurement.frequency.SETTLE_DELAY', 0):
            results, status = engine.run(['0.01', '0.1'])
        self.assertEqual(status, 'error')
        self.assertEqual(len(results[TEST_FREQ]), 1)
        self.assertEqual(results[TEST_FREQ][0].status, 'error')
        self.assertEqual(results[TEST_FREQ][0].reference_frequency_hz, 10000.0)
        self.assertEqual(results[TEST_FREQ][0].comment,
                         "Некорректный разбор частоты ПЧ: could not convert string to float: 'invalid'")
        self.assertIn('OUTP OFF', engine.gen.session.commands)
        self.assertTrue(engine.diagnostics)

    def test_init_failure_attempts_off_and_logging_failure_is_ignored(self):
        gen = FakeGenerator(fail_write='OUTP ON')
        engine = MeasurementEngine(gen, FakeMeter(), [TEST_FREQ], log_callback=lambda _: (_ for _ in ()).throw(RuntimeError()))
        results, status = engine.run(['0.01'])
        self.assertEqual(status, 'error')
        self.assertIn('OUTP OFF', gen.commands)
        self.assertIsNone(results[TEST_FREQ][0] if results[TEST_FREQ] else None)

    def test_shutdown_can_be_retried(self):
        gen = FakeGenerator(fail_write='OUTP OFF')
        engine = self.engine(gen=gen)
        with patch('stand_verification.measurement.frequency.SETTLE_DELAY', 0):
            _, status = engine.run(['0.01'])
        self.assertEqual(status, 'error')
        self.assertFalse(engine.shutdown_success)
        gen.fail_write = None
        self.assertEqual(engine._deinit_generator(), (True, ''))
        self.assertTrue(engine.shutdown_success)

    def test_stop_before_start_does_not_enable_generator(self):
        import threading
        stop = threading.Event()
        stop.set()
        engine = self.engine(stop_event=stop)
        _, status = engine.run(['0.01'])
        self.assertEqual(status, 'stopped')
        self.assertEqual(engine.gen.session.commands, [])
        self.assertIsNone(engine.shutdown_success)

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'openpyxl is not installed')
    def test_export_has_summary_and_error_status(self):
        engine = self.engine()
        engine.results[TEST_FREQ] = [FrequencyResult(
            reference_frequency_hz=10000, measured_frequency_hz=None,
            error_hz=None, tolerance_hz=None, if_frequency_khz=None,
            status='error', comment='meter failure',
        )]
        engine.final_status = 'error'
        engine._planned = {TEST_FREQ: 1}
        engine.diagnostics = ['meter failure']
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.xlsx'
            self.assertIs(export_excel(engine.snapshot(), path, engine._log), True)
            from openpyxl import load_workbook
            workbook = load_workbook(path)
            sheet = workbook['Частота']
            self.assertEqual(sheet['F2'].value, 'Ошибка')
            self.assertIn('Сводка', workbook.sheetnames)
            self.assertIn('Диагностика', workbook.sheetnames)

    def test_export_failure_propagates(self):
        engine = self.engine()
        with self.assertRaises(Exception):
            export_excel(engine.snapshot(), '/nonexistent/path/report.xlsx', engine._log)

    def test_level_tolerance_boundary_and_evaluation(self):
        cases = [(1999, 1.5, 1.0, 'fail'),
                 (2000, 1.5, 2.0, 'warn'),
                 (2500, -2.0, 2.0, 'warn'),
                 (3000, 2.01, 2.0, 'fail')]
        for frequency, measured, tolerance, expected in cases:
            with self.subTest(frequency=frequency, measured=measured):
                engine = self.engine(sk3=FakeMeter([str(measured).encode()]))
                with patch('stand_verification.measurement.level.SETTLE_DELAY', 0):
                    row = engine._measure_level(str(frequency), 0)
                self.assertEqual(row.tolerance_db, tolerance)
                self.assertEqual(row.status, expected)

    def test_level_error_keeps_frequency_dependent_tolerance(self):
        engine = self.engine(gen=FakeGenerator(fail_write='POW 0'))
        row = engine._measure_level('2000', 0)
        self.assertEqual(row, LevelResult(
            frequency_mhz=2000.0, set_level_dbm=0,
            measured_level_dbm=None, error_db=None, tolerance_db=2.0,
            status='error', comment='Ошибка установки уровня 0 dBm: write failed: POW 0',
        ))


if __name__ == '__main__':
    unittest.main()
